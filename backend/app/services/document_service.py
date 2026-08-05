"""
FinPilot AI — Document Processing Service
==========================================
Handles the full document ingestion pipeline:
1. File validation and storage
2. Content extraction (text, tables, images)
3. OCR for scanned documents
4. Semantic chunking
5. Metadata extraction
6. Embedding generation
7. Vector storage
"""

import hashlib
import io
import logging
import os
import time
import uuid
from pathlib import Path
from typing import Optional

from app.core.config import settings
from app.services.vector_service import VectorService

logger = logging.getLogger(__name__)


class DocumentProcessingService:
    """
    End-to-end document processing pipeline for financial documents.
    Supports PDF, DOCX, PPTX, Excel, CSV, TXT, and images.
    """

    def __init__(self):
        self.vector_service = VectorService()
        self.upload_dir = Path(settings.UPLOAD_DIR)
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    def validate_file(self, filename: str, file_size: int, content_type: str) -> None:
        """
        Validate uploaded file before processing.

        Raises:
            ValueError: If file type or size is invalid
        """
        # Check file size
        if file_size > settings.max_file_size_bytes:
            raise ValueError(
                f"File size {file_size / 1024 / 1024:.1f}MB exceeds "
                f"maximum allowed {settings.MAX_FILE_SIZE_MB}MB"
            )

        # Check file extension
        ext = Path(filename).suffix.lower().lstrip(".")
        if ext not in settings.ALLOWED_FILE_TYPES:
            raise ValueError(
                f"File type '.{ext}' is not supported. "
                f"Allowed types: {', '.join(settings.ALLOWED_FILE_TYPES)}"
            )

    async def save_file(self, file_content: bytes, original_filename: str, user_id: str) -> tuple[str, str, str]:
        """
        Save uploaded file to disk.

        Returns:
            Tuple of (saved_filename, file_path, sha256_hash)
        """
        # Generate unique filename
        ext = Path(original_filename).suffix.lower()
        unique_filename = f"{user_id}_{uuid.uuid4().hex}{ext}"
        file_path = self.upload_dir / unique_filename

        # Write file
        with open(file_path, "wb") as f:
            f.write(file_content)

        # Compute SHA256 hash
        sha256 = hashlib.sha256(file_content).hexdigest()

        return unique_filename, str(file_path), sha256

    async def extract_text_from_pdf(self, file_path: str) -> tuple[str, int, list[dict]]:
        """
        Extract text and tables from PDF.

        Returns:
            Tuple of (full_text, page_count, extracted_tables)
        """
        try:
            import pdfplumber

            full_text = ""
            tables = []
            page_count = 0

            with pdfplumber.open(file_path) as pdf:
                page_count = len(pdf.pages)
                for page_num, page in enumerate(pdf.pages):
                    # Extract text
                    page_text = page.extract_text() or ""
                    if page_text:
                        full_text += f"\n\n[Page {page_num + 1}]\n{page_text}"

                    # Extract tables
                    if settings.TABLE_EXTRACTION_ENABLED:
                        page_tables = page.extract_tables()
                        for table in page_tables:
                            if table:
                                tables.append({
                                    "page": page_num + 1,
                                    "data": table,
                                    "rows": len(table),
                                    "cols": len(table[0]) if table else 0,
                                })

            # Apply OCR if text extraction yielded little content
            if settings.OCR_ENABLED and len(full_text.strip()) < 100:
                logger.info(f"Low text content detected, applying OCR for {file_path}")
                full_text = await self._apply_ocr(file_path)

            return full_text.strip(), page_count, tables

        except ImportError:
            logger.warning("pdfplumber not available, using fallback PDF extraction")
            return await self._extract_text_fallback(file_path)
        except Exception as e:
            logger.error(f"PDF extraction failed: {e}")
            raise

    async def extract_text_from_docx(self, file_path: str) -> tuple[str, int, list[dict]]:
        """Extract text and tables from DOCX files."""
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)

            # Extract tables from DOCX
            tables = []
            for i, table in enumerate(doc.tables):
                table_data = []
                for row in table.rows:
                    table_data.append([cell.text for cell in row.cells])
                tables.append({
                    "page": 1,
                    "data": table_data,
                    "rows": len(table_data),
                    "cols": len(table_data[0]) if table_data else 0,
                })

            return full_text, 1, tables

        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            raise

    async def extract_text_from_excel(self, file_path: str) -> tuple[str, int, list[dict]]:
        """Extract data from Excel/CSV files."""
        try:
            import pandas as pd

            ext = Path(file_path).suffix.lower()
            if ext == ".csv":
                df_dict = {"Sheet1": pd.read_csv(file_path)}
            else:
                df_dict = pd.read_excel(file_path, sheet_name=None)

            full_text = ""
            tables = []

            for sheet_name, df in df_dict.items():
                # Convert to text representation
                full_text += f"\n\n[Sheet: {sheet_name}]\n"
                full_text += df.to_string(index=False)

                # Store as structured table
                tables.append({
                    "sheet": sheet_name,
                    "data": [df.columns.tolist()] + df.fillna("").values.tolist(),
                    "rows": len(df),
                    "cols": len(df.columns),
                    "summary": {
                        "shape": df.shape,
                        "dtypes": df.dtypes.astype(str).to_dict(),
                        "numeric_cols": df.select_dtypes(include="number").columns.tolist(),
                    },
                })

            return full_text.strip(), 1, tables

        except Exception as e:
            logger.error(f"Excel/CSV extraction failed: {e}")
            raise

    async def _apply_ocr(self, file_path: str) -> str:
        """Apply OCR using pytesseract for scanned documents."""
        try:
            import pytesseract
            from PIL import Image
            from pdf2image import convert_from_path

            ext = Path(file_path).suffix.lower()
            if ext == ".pdf":
                pages = convert_from_path(file_path, dpi=300)
                text_parts = []
                for i, page in enumerate(pages):
                    page_text = pytesseract.image_to_string(
                        page, config="--psm 6 -l eng"
                    )
                    text_parts.append(f"[Page {i + 1}]\n{page_text}")
                return "\n\n".join(text_parts)
            else:
                img = Image.open(file_path)
                return pytesseract.image_to_string(img)
        except ImportError:
            logger.warning("OCR libraries not available (pytesseract/pdf2image)")
            return ""
        except Exception as e:
            logger.warning(f"OCR failed: {e}")
            return ""

    async def _extract_text_fallback(self, file_path: str) -> tuple[str, int, list]:
        """Simple text extraction fallback."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(), 1, []
        except Exception:
            return "", 0, []

    def semantic_chunk(
        self,
        text: str,
        chunk_size: int = None,
        chunk_overlap: int = None,
    ) -> list[str]:
        """
        Split text into overlapping chunks using recursive character text splitting.
        Uses sentence boundaries where possible for semantic coherence.

        Args:
            text: Full document text
            chunk_size: Max characters per chunk
            chunk_overlap: Overlap between chunks

        Returns:
            List of text chunks
        """
        chunk_size = chunk_size or settings.CHUNK_SIZE
        chunk_overlap = chunk_overlap or settings.CHUNK_OVERLAP

        try:
            from langchain_text_splitters import RecursiveCharacterTextSplitter

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                separators=["\n\n", "\n", ". ", "! ", "? ", " ", ""],
                keep_separator=True,
            )
            chunks = splitter.split_text(text)
            return chunks[:settings.MAX_CHUNKS_PER_DOCUMENT]
        except ImportError:
            # Fallback manual chunking
            chunks = []
            start = 0
            while start < len(text):
                end = min(start + chunk_size, len(text))
                chunks.append(text[start:end])
                start += chunk_size - chunk_overlap
            return chunks[:settings.MAX_CHUNKS_PER_DOCUMENT]

    async def process_document(
        self,
        document_id: str,
        file_path: str,
        doc_type: str,
        user_id: str,
        metadata: Optional[dict] = None,
    ) -> dict:
        """
        Full document processing pipeline.
        Extracts text → chunks → embeds → stores in ChromaDB.

        Returns:
            Processing result with statistics
        """
        start_time = time.perf_counter()
        result = {
            "document_id": document_id,
            "status": "processing",
            "chunks_count": 0,
            "tables_count": 0,
            "page_count": 0,
            "ocr_applied": False,
            "processing_time_seconds": 0,
            "error": None,
        }

        try:
            # ── Step 1: Extract text based on file type ────────────────────────
            ext = Path(file_path).suffix.lower()

            if ext == ".pdf":
                text, page_count, tables = await self.extract_text_from_pdf(file_path)
                result["ocr_applied"] = len(text.strip()) < 100
            elif ext in (".docx", ".doc"):
                text, page_count, tables = await self.extract_text_from_docx(file_path)
            elif ext in (".xlsx", ".xls", ".csv"):
                text, page_count, tables = await self.extract_text_from_excel(file_path)
            elif ext in (".pptx", ".ppt"):
                text, page_count, tables = await self._extract_pptx(file_path)
            elif ext in (".txt", ".md"):
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
                page_count, tables = 1, []
            elif ext in (".png", ".jpg", ".jpeg", ".tiff"):
                text = await self._apply_ocr(file_path)
                page_count, tables = 1, []
                result["ocr_applied"] = True
            else:
                text, page_count, tables = await self._extract_text_fallback(file_path)

            result["page_count"] = page_count
            result["tables_count"] = len(tables)

            if not text.strip():
                raise ValueError("No text content could be extracted from the document")

            # ── Step 2: Semantic chunking ──────────────────────────────────────
            chunks = self.semantic_chunk(text)
            result["chunks_count"] = len(chunks)

            # ── Step 3: Prepare chunk metadata ────────────────────────────────
            chunk_metadatas = [
                {
                    "document_id": document_id,
                    "user_id": user_id,
                    "chunk_index": i,
                    "doc_type": doc_type,
                    "file_path": file_path,
                    **(metadata or {}),
                }
                for i in range(len(chunks))
            ]

            # ── Step 4: Store in ChromaDB ──────────────────────────────────────
            vector_ids = await self.vector_service.add_documents(
                document_id=document_id,
                chunks=chunks,
                metadatas=chunk_metadatas,
            )

            result["status"] = "processed"
            result["vector_ids"] = vector_ids
            result["raw_text"] = text[:5000]  # Store first 5000 chars in DB

        except Exception as e:
            logger.error(f"Document processing failed for {document_id}: {e}")
            result["status"] = "failed"
            result["error"] = str(e)

        finally:
            result["processing_time_seconds"] = round(
                time.perf_counter() - start_time, 2
            )

        return result

    async def _extract_pptx(self, file_path: str) -> tuple[str, int, list]:
        """Extract text from PowerPoint presentations."""
        try:
            from pptx import Presentation

            prs = Presentation(file_path)
            text_parts = []

            for i, slide in enumerate(prs.slides):
                slide_text = f"[Slide {i + 1}]"
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text += f"\n{shape.text}"
                text_parts.append(slide_text)

            return "\n\n".join(text_parts), len(prs.slides), []
        except Exception as e:
            logger.error(f"PPTX extraction failed: {e}")
            return "", 0, []
